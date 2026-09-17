"""
Comprehensive test script for multi-format RAG with ChromaDB.

Tests:
1. Text extraction for all supported formats (PDF, DOCX, PPTX, TXT, CSV, XLSX)
2. Text chunking logic
3. ChromaDB ingestion and semantic search
4. End-to-end: extract → chunk → ingest → search
5. Existing PDF functionality is preserved

Usage:
    cd ai-chatbot
    python test_multi_format_rag.py
"""
import os
import sys
import shutil
import tempfile
import asyncio

# ── Ensure project root is in path ──────────────────────────────────────────
sys.path.insert(0, os.path.dirname(__file__))

TEMP_DIR = os.path.join(os.path.dirname(__file__), "_test_temp")
TEST_USER_ID = "test_user_multiformat_rag"

passed = 0
failed = 0
errors = []


def _report(test_name: str, success: bool, detail: str = ""):
    global passed, failed
    if success:
        passed += 1
        print(f"  ✅ {test_name}")
    else:
        failed += 1
        errors.append((test_name, detail))
        print(f"  ❌ {test_name}: {detail}")


# ═══════════════════════════════════════════════════════════════════════════
# PHASE 1: Create sample files
# ═══════════════════════════════════════════════════════════════════════════

def create_sample_files():
    """Create sample files for each supported format."""
    os.makedirs(TEMP_DIR, exist_ok=True)
    files = {}

    # ── TXT ───────────────────────────────────────────────────────────────
    txt_path = os.path.join(TEMP_DIR, "sample.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("Artificial intelligence is transforming healthcare.\n\n"
                "Machine learning models can predict disease outcomes.\n\n"
                "Deep learning has revolutionized medical imaging analysis.")
    files["txt"] = txt_path

    # ── CSV ───────────────────────────────────────────────────────────────
    csv_path = os.path.join(TEMP_DIR, "sample.csv")
    with open(csv_path, "w", encoding="utf-8") as f:
        f.write("Name,Role,Department\n")
        f.write("Alice,Engineer,AI Research\n")
        f.write("Bob,Manager,Product Development\n")
        f.write("Charlie,Designer,User Experience\n")
    files["csv"] = csv_path

    # ── DOCX ──────────────────────────────────────────────────────────────
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("Climate Change Report", level=1)
        doc.add_paragraph("Global temperatures have risen by 1.1°C since pre-industrial times.")
        doc.add_paragraph("Renewable energy adoption is accelerating worldwide.")
        doc.add_paragraph("Carbon capture technology shows promising results in pilot programs.")
        docx_path = os.path.join(TEMP_DIR, "sample.docx")
        doc.save(docx_path)
        files["docx"] = docx_path
    except ImportError:
        print("  ⚠️  python-docx not installed, skipping DOCX test file creation")

    # ── PPTX ──────────────────────────────────────────────────────────────
    try:
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title + Content
        slide.shapes.title.text = "Quantum Computing Overview"
        slide.placeholders[1].text = (
            "Quantum computers use qubits instead of classical bits.\n"
            "Quantum entanglement enables parallel processing.\n"
            "Applications include cryptography and drug discovery."
        )
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Future Prospects"
        slide2.placeholders[1].text = (
            "Error correction is the main challenge.\n"
            "Hybrid quantum-classical algorithms are emerging."
        )
        pptx_path = os.path.join(TEMP_DIR, "sample.pptx")
        prs.save(pptx_path)
        files["pptx"] = pptx_path
    except ImportError:
        print("  ⚠️  python-pptx not installed, skipping PPTX test file creation")

    # ── XLSX ──────────────────────────────────────────────────────────────
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Data"
        ws.append(["Product", "Revenue", "Quarter"])
        ws.append(["Widget Alpha", 150000, "Q1"])
        ws.append(["Widget Beta", 230000, "Q2"])
        ws.append(["Widget Gamma", 180000, "Q3"])
        xlsx_path = os.path.join(TEMP_DIR, "sample.xlsx")
        wb.save(xlsx_path)
        files["xlsx"] = xlsx_path
    except ImportError:
        print("  ⚠️  openpyxl not installed, skipping XLSX test file creation")

    # ── PDF ────────────────────────────────────────────────────────────────
    try:
        import fitz  # PyMuPDF
        pdf_path = os.path.join(TEMP_DIR, "sample.pdf")
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Neural Networks and Deep Learning\n\n"
                         "Neural networks are inspired by biological neurons.\n"
                         "Backpropagation is the key training algorithm.\n"
                         "Convolutional neural networks excel at image recognition.",
                         fontsize=12)
        doc.save(pdf_path)
        doc.close()
        files["pdf"] = pdf_path
    except ImportError:
        print("  ⚠️  PyMuPDF not installed, skipping PDF test file creation")

    return files


# ═══════════════════════════════════════════════════════════════════════════
# PHASE 2: Test Text Extraction
# ═══════════════════════════════════════════════════════════════════════════

async def test_text_extraction(files: dict):
    """Test that text_extractor works for each format."""
    print("\n── Phase 2: Text Extraction ──")
    from app.services.text_extractor import extract_text, is_supported, detect_file_type

    extracted = {}

    for fmt, path in files.items():
        filename = os.path.basename(path)

        # Test detection
        detected = detect_file_type(filename)
        _report(f"{fmt.upper()} type detection", detected == fmt,
                f"Expected '{fmt}', got '{detected}'")

        # Test is_supported
        supported = is_supported(filename)
        _report(f"{fmt.upper()} is_supported", supported, "Not recognized as supported")

        # Test extraction
        try:
            text, ftype = await extract_text(path, filename)
            has_content = len(text.strip()) > 10
            _report(f"{fmt.upper()} extraction ({len(text)} chars)",
                    has_content, f"Extracted text too short: '{text[:50]}'")
            if has_content:
                extracted[fmt] = text
        except Exception as e:
            _report(f"{fmt.upper()} extraction", False, str(e))

    return extracted


# ═══════════════════════════════════════════════════════════════════════════
# PHASE 3: Test Chunking
# ═══════════════════════════════════════════════════════════════════════════

def test_chunking(extracted: dict):
    """Test the text chunking logic."""
    print("\n── Phase 3: Text Chunking ──")
    from app.services.vector_store import chunk_text, CHUNK_SIZE

    for fmt, text in extracted.items():
        chunks = chunk_text(text)
        has_chunks = len(chunks) > 0
        _report(f"{fmt.upper()} chunking ({len(chunks)} chunks)", has_chunks,
                "No chunks produced")

        # Verify no chunk is excessively long
        if chunks:
            max_len = max(len(c) for c in chunks)
            reasonable = max_len < CHUNK_SIZE * 2
            _report(f"{fmt.upper()} chunk size check (max={max_len})", reasonable,
                    f"Chunk too long: {max_len} chars")

    # Edge cases
    _report("Empty text chunking", len(chunk_text("")) == 0, "Should return empty list")
    _report("Whitespace-only chunking", len(chunk_text("   \n\n  ")) == 0, "Should return empty list")
    _report("Short text chunking", len(chunk_text("Hello world")) == 0,
            "Text below MIN_CHUNK_LENGTH should be skipped")


# ═══════════════════════════════════════════════════════════════════════════
# PHASE 4: Test ChromaDB Ingestion & Search
# ═══════════════════════════════════════════════════════════════════════════

def test_chromadb(extracted: dict):
    """Test ChromaDB ingestion and semantic search."""
    print("\n── Phase 4: ChromaDB Ingestion & Search ──")
    from app.services.vector_store import ingest_document, search_documents, delete_document

    total_chunks = 0

    # Ingest all extracted documents
    for fmt, text in extracted.items():
        filename = f"sample.{fmt}"
        try:
            count = ingest_document(
                user_id=TEST_USER_ID,
                filename=filename,
                text_content=text,
                file_type=fmt,
            )
            total_chunks += count
            _report(f"{fmt.upper()} ingestion ({count} chunks)", count > 0,
                    "No chunks ingested")
        except Exception as e:
            _report(f"{fmt.upper()} ingestion", False, str(e))

    _report(f"Total chunks ingested: {total_chunks}", total_chunks > 0, "Nothing was ingested")

    # ── Semantic Search Tests ────────────────────────────────────────────

    # Test 1: Search for healthcare/AI content (should match TXT)
    if "txt" in extracted:
        results = search_documents(TEST_USER_ID, "healthcare artificial intelligence")
        found_txt = any(r["filename"] == "sample.txt" for r in results)
        _report("Semantic search: AI healthcare → TXT", found_txt,
                f"Expected sample.txt in results, got: {[r['filename'] for r in results]}")

    # Test 2: Search for climate content (should match DOCX)
    if "docx" in extracted:
        results = search_documents(TEST_USER_ID, "climate change temperature renewable energy")
        found_docx = any(r["filename"] == "sample.docx" for r in results)
        _report("Semantic search: climate → DOCX", found_docx,
                f"Expected sample.docx in results, got: {[r['filename'] for r in results]}")

    # Test 3: Search for quantum content (should match PPTX)
    if "pptx" in extracted:
        results = search_documents(TEST_USER_ID, "quantum computing qubits")
        found_pptx = any(r["filename"] == "sample.pptx" for r in results)
        _report("Semantic search: quantum → PPTX", found_pptx,
                f"Expected sample.pptx in results, got: {[r['filename'] for r in results]}")

    # Test 4: Search for sales data (should match XLSX)
    if "xlsx" in extracted:
        results = search_documents(TEST_USER_ID, "sales revenue product widget")
        found_xlsx = any(r["filename"] == "sample.xlsx" for r in results)
        _report("Semantic search: sales → XLSX", found_xlsx,
                f"Expected sample.xlsx in results, got: {[r['filename'] for r in results]}")

    # Test 5: Search for neural networks (should match PDF)
    if "pdf" in extracted:
        results = search_documents(TEST_USER_ID, "neural networks deep learning backpropagation")
        found_pdf = any(r["filename"] == "sample.pdf" for r in results)
        _report("Semantic search: neural networks → PDF", found_pdf,
                f"Expected sample.pdf in results, got: {[r['filename'] for r in results]}")

    # Test 6: Results have proper metadata
    if total_chunks > 0:
        results = search_documents(TEST_USER_ID, "test query", top_k=3)
        if results:
            r = results[0]
            has_meta = all(k in r for k in ["text", "filename", "file_type", "score"])
            _report("Search result metadata fields", has_meta,
                    f"Missing fields. Keys: {list(r.keys())}")
            has_score = 0 <= r["score"] <= 1
            _report("Search result score in [0,1]", has_score,
                    f"Score out of range: {r['score']}")

    # Test 7: Delete document
    if "txt" in extracted:
        deleted = delete_document(TEST_USER_ID, "sample.txt")
        _report("Delete document", deleted, "Delete returned False")

        # Verify it's gone from search
        results = search_documents(TEST_USER_ID, "healthcare artificial intelligence")
        not_found = not any(r["filename"] == "sample.txt" for r in results)
        _report("Deleted document not in search results", not_found,
                "sample.txt still appears after deletion")


# ═══════════════════════════════════════════════════════════════════════════
# PHASE 5: PDF-specific backward compatibility
# ═══════════════════════════════════════════════════════════════════════════

async def test_pdf_backward_compat(files: dict):
    """Ensure existing PDF pipeline still works identically."""
    print("\n── Phase 5: PDF Backward Compatibility ──")

    if "pdf" not in files:
        print("  ⚠️  No PDF test file, skipping backward compat tests")
        return

    from app.services.text_extractor import extract_text

    pdf_path = files["pdf"]
    text, ftype = await extract_text(pdf_path, "sample.pdf")
    _report("PDF extraction returns text", len(text.strip()) > 10, "Empty text")
    _report("PDF file_type is 'pdf'", ftype == "pdf", f"Got '{ftype}'")

    # Verify .txt sidecar convention
    txt_sidecar = pdf_path + ".txt"
    with open(txt_sidecar, "w", encoding="utf-8") as f:
        f.write(text)
    _report("PDF .txt sidecar writable", os.path.exists(txt_sidecar), "Failed to write sidecar")

    # Verify keyword search on sidecar still works
    with open(txt_sidecar, "r", encoding="utf-8") as f:
        content = f.read()
    chunks = [p.strip() for p in content.split("\n\n") if p.strip()]
    _report("PDF sidecar has chunks", len(chunks) > 0, "No chunks from sidecar")


# ═══════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════

async def main():
    global passed, failed, errors

    print("=" * 65)
    print("  Multi-Format RAG + ChromaDB — Comprehensive Test Suite")
    print("=" * 65)

    # Phase 1: Create sample files
    print("\n── Phase 1: Creating Sample Files ──")
    files = create_sample_files()
    print(f"  Created {len(files)} sample files: {list(files.keys())}")

    # Phase 2: Text Extraction
    extracted = await test_text_extraction(files)

    # Phase 3: Chunking
    test_chunking(extracted)

    # Phase 4: ChromaDB
    test_chromadb(extracted)

    # Phase 5: PDF backward compat
    await test_pdf_backward_compat(files)

    # ── Summary ──────────────────────────────────────────────────────────
    print("\n" + "=" * 65)
    total = passed + failed
    print(f"  Results: {passed}/{total} passed, {failed} failed")
    if errors:
        print("\n  Failed tests:")
        for name, detail in errors:
            print(f"    ❌ {name}: {detail}")
    print("=" * 65)

    # Cleanup
    try:
        shutil.rmtree(TEMP_DIR, ignore_errors=True)
        # Clean up test user ChromaDB data
        from app.services.vector_store import _get_chroma_client
        import re
        client = _get_chroma_client()
        safe_name = "user_" + re.sub(r"[^a-zA-Z0-9_]", "_", TEST_USER_ID)[:55]
        try:
            client.delete_collection(safe_name)
        except Exception:
            pass
    except Exception:
        pass

    return failed == 0


if __name__ == "__main__":
    success = asyncio.run(main())
    sys.exit(0 if success else 1)
